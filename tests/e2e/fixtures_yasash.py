"""E2E fiksturalarini BIR MARTA yasagan skript (hujjat sifatida saqlanadi).

Fayllar MUZLATILGAN (`tests/e2e/fixtures/`): ssenariy baytma-bayt bir xil
kirishlarda ishlashi shart, openpyxl esa har saqlashda vaqtni yozadi — qayta
yasalsa sha256 lar va natijalar o'zgaradi. Qayta yasash kerak bo'lsa,
`tests/e2e/kutilgan.json` ni ham yangilang (E2E_YANGILA=1).

`jam_*` — jamoaning o'z namunalari (monolitdagi `services/back/tender/namuna`),
qolganlari shu yerda sintez qilinadi: to'rt hujjat turi (jamlanma, narxlar,
resurs, loyiha), har birining shabloni va darvozalarni yurgizadigan
variantlari, notanish shakl va buzuq fayllar.

    python tests/e2e/fixtures_yasash.py CHIQISH_PAPKA NAMUNA_PAPKA
"""
import datetime
import os
import shutil
import sys

import openpyxl

OUT = sys.argv[1]
NAMUNA = sys.argv[2]
FIXED = datetime.datetime(2026, 9, 1, 12, 0, 0)


def save(wb, name):
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    wb.save(os.path.join(OUT, name))


def book(rows, title="Лист1", extra=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    for t, rr in (extra or []):
        w2 = wb.create_sheet(t)
        for r in rr:
            w2.append(r)
    return wb


os.makedirs(OUT, exist_ok=True)

# ── jamlanma (excel1): the team's own samples, verdicts documented in namuna/README.md
for f in os.listdir(NAMUNA):
    if f.endswith(".xlsx"):
        shutil.copy2(os.path.join(NAMUNA, f), os.path.join(OUT, "jam_" + f))

# ── narxlar (excel3) ────────────────────────────────────────────────────────
NARX_HEAD = ["№", "Компонент номи", "Ўлчов бирлиги", "Миқдори", "Бирлик учун нарх", "Жами"]
ITEMS = [("Цемент М400", "т", 12), ("Қум", "м3", 40), ("Шағал", "м3", 35),
         ("Арматура А500", "т", 6), ("Ғишт", "минг дона", 18), ("Бўёқ", "кг", 120),
         ("Ойна", "м2", 64), ("Кабель ВВГ", "м", 900)]


def narx(prices=None, head=NARX_HEAD, drop=None, add=None, errors=False):
    rows = [["Нархлар жадвали"], [], head]
    for i, (n, u, q) in enumerate(ITEMS, 1):
        p = prices[i - 1] if prices else None
        row = [i, n, u, q, p, (p * q if isinstance(p, (int, float)) else None)]
        if errors and i in (2, 5):
            row[4] = "#VALUE!"
        if add:
            row.append(add[i - 1])
        if drop is not None:
            row.pop(drop)
        rows.append(row)
    if add:
        rows[2] = head + ["Таклиф нархи"]
    if drop is not None:
        rows[2] = [h for j, h in enumerate(rows[2]) if j != drop]
    return book(rows)


FULL = [1250000, 90000, 110000, 9800000, 1450000, 42000, 185000, 21000]
save(narx(), "narx_shablon.xlsx")
save(narx(FULL), "narx_toldirilgan.xlsx")
save(narx([None] * 7 + [21000]), "narx_siyrak.xlsx")
save(narx(FULL, drop=4), "narx_ustun_ochirilgan.xlsx")
save(narx(None, add=FULL), "narx_qoshimcha_ustun.xlsx")
save(narx(FULL, errors=True), "narx_xato_qiymat.xlsx")
wb = openpyxl.load_workbook(os.path.join(OUT, "narx_shablon.xlsx"))
wb.save(os.path.join(OUT, "narx_qayta_saqlangan.xlsx"))          # same cells, new bytes

# ── resurs (excel2): multi-sheet local estimate ─────────────────────────────
RES_HEAD = ["№ п/п", "Шифр номера нормативов", "Наименование работ и затрат",
            "Ед.изм", "Кол-во", "Стоимость"]
WORKS = [("Е1-1", "Разработка грунта", "м3", 120), ("Е6-1", "Бетон фундамента", "м3", 45),
         ("Е8-3", "Кладка стен", "м3", 80), ("Е12-7", "Кровля", "м2", 300),
         ("Е15-2", "Штукатурка", "м2", 650), ("Е26-1", "Утепление", "м2", 280)]


def resurs(prices=None, add_col=None, second_prices=None):
    def sheet(ps):
        head = RES_HEAD + ([add_col] if add_col else [])
        rows = [["Локальная ресурсная ведомость"], [], head]
        for i, (code, name, unit, qty) in enumerate(WORKS, 1):
            p = ps[i - 1] if ps else None
            r = [i, code, name, unit, qty, None if add_col else p]
            if add_col:
                r.append(p)
            rows.append(r)
        return rows
    return book(sheet(prices), title="1",
                extra=[("2", sheet(second_prices))])


RP = [520000, 1450000, 980000, 210000, 65000, 88000]
save(resurs(), "res_shablon.xlsx")
save(resurs(RP, second_prices=RP), "res_toldirilgan.xlsx")
save(resurs(RP), "res_bir_varaq.xlsx")                            # second sheet left empty
save(resurs(RP, add_col="Предложение претендента", second_prices=RP), "res_qoshimcha_ustun.xlsx")
save(resurs([None] * 6, second_prices=[None] * 5 + [88000]), "res_siyrak.xlsx")

# ── loyiha (loyiha_excel): «5-ILOVA» / SMETA HISOBI ─────────────────────────
LOY_HEAD = ["T/r", "Ish turlari", "Asos", "Narxning kelib chiqishi", "Qiymat (ming so'm)"]
LOY = [("Loyiha-qidiruv ishlari", "QMQ 1.03.01", "Hisob"),
       ("Arxitektura yechimlari", "QMQ 2.08.02", "Hisob"),
       ("Konstruktiv yechimlar", "QMQ 2.03.01", "Hisob"),
       ("Muhandislik tarmoqlari", "QMQ 3.05.01", "Hisob"),
       ("Smeta hujjatlari", "QMQ 4.02.01", "Hisob")]


def loyiha(values=None, drop_price=False):
    head = LOY_HEAD[:-1] if drop_price else LOY_HEAD
    rows = [["5-ILOVA"], ["SMETA HISOBI"], [], head]
    for i, (n, a, k) in enumerate(LOY, 1):
        r = [i, n, a, k]
        if not drop_price:
            r.append(values[i - 1] if values else None)
        rows.append(r)
    rows.append(["", "Jami", "", "", (sum(v for v in values if v) or None) if values and not drop_price else None])
    return book(rows, title="5-ILOVA")


LV = [12000, 30500, 27400, 18800, 6300]
save(loyiha(), "loy_shablon.xlsx")
save(loyiha(LV), "loy_toldirilgan.xlsx")
save(loyiha([None] * 5), "loy_bosh.xlsx")
save(loyiha(LV, drop_price=True), "loy_ustun_ochirilgan.xlsx")

# ── an unrecognisable layout: no role signature anywhere (type decides) ─────
ODD = [["Хужжат"], ["Колонка A", "Колонка B", "Колонка C"],
       ["x1", 3, None], ["x2", 5, None], ["x3", 8, None], ["x4", 13, None]]
save(book(ODD), "odd_shablon.xlsx")
save(book([r if i < 2 else r[:2] + [r[1] * 1000] for i, r in enumerate(ODD)]), "odd_toldirilgan.xlsx")

# ── broken files ─────────────────────────────────────────────────────────────
with open(os.path.join(OUT, "buzuq_html.xlsx"), "wb") as f:           # an HTML error page
    f.write(b"<html><body>Access denied</body></html>")
with open(os.path.join(OUT, "buzuq_zip.xlsx"), "wb") as f:            # zip magic, not a workbook
    f.write(b"PK\x03\x04" + b"\x00" * 60 + b"not really a zip archive at all" * 8)
with open(os.path.join(OUT, "bosh_fayl.xlsx"), "wb") as f:            # zero bytes
    pass

print("\n".join(sorted(os.listdir(OUT))))
