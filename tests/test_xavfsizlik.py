# -*- coding: utf-8 -*-
"""
Xavfsizlik va deploy tayyorligi testlari (baza KERAK EMAS).

    pytest tests/ -v
"""
import os
import sys
import time
import shutil
import zipfile
import tempfile
import threading
import subprocess
from http.server import BaseHTTPRequestHandler, HTTPServer

import pytest

ILDIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


sys.path.insert(0, ILDIZ)

from tender_engine.reader import read_file, ExcelTooLargeError   # noqa: E402
from app import download as common                               # noqa: E402

# Ishlab chiqarishga chiqadigan fayllar — Dockerfile ham AYNAN shularni ko'chiradi
# (paketlar butunicha, kirish skriptlari ro'yxat bilan).
PROD_FAYLLAR = [
    "main.py", "jobs_worker.py", "yuboruvchi.py", "api_server.py",
    "korish.py", "kuzatuv.py", "ishga_tushir.py", "requirements.txt",
]
PROD_PAKETLAR = ["app", "tender_engine"]


# ═══════════════════════════════════════════════════════════════════════════
# Excel merge-bomba (blocker 2)
# ═══════════════════════════════════════════════════════════════════════════

def _merge_bomba(yol, ref="A1:XFD1048576"):
    """Kichkina, lekin ulkan merge diapazonli haqiqiy .xlsx yasaydi."""
    sh = ('<?xml version="1.0"?><worksheet xmlns="http://schemas.openxmlformats.org/'
          'spreadsheetml/2006/main"><dimension ref="A1:D10"/><sheetData>'
          '<row r="1"><c r="A1" t="inlineStr"><is><t>x</t></is></c></row>'
          f'</sheetData><mergeCells count="1"><mergeCell ref="{ref}"/></mergeCells></worksheet>')
    wb = ('<?xml version="1.0"?><workbook xmlns="http://schemas.openxmlformats.org/'
          'spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/'
          'officeDocument/2006/relationships"><sheets><sheet name="Sheet1" '
          'sheetId="1" r:id="rId1"/></sheets></workbook>')
    rl = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/'
          'package/2006/relationships"><Relationship Id="rId1" Type="http://'
          'schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
          'Target="xl/workbook.xml"/></Relationships>')
    wr = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/'
          'package/2006/relationships"><Relationship Id="rId1" Type="http://'
          'schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
          'Target="worksheets/sheet1.xml"/></Relationships>')
    ct = ('<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/'
          'package/2006/content-types"><Default Extension="rels" ContentType='
          '"application/vnd.openxmlformats-package.relationships+xml"/><Default '
          'Extension="xml" ContentType="application/xml"/><Override PartName='
          '"/xl/workbook.xml" ContentType="application/vnd.openxmlformats-'
          'officedocument.spreadsheetml.sheet.main+xml"/><Override PartName='
          '"/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-'
          'officedocument.spreadsheetml.worksheet+xml"/></Types>')
    with zipfile.ZipFile(yol, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rl)
        z.writestr("xl/workbook.xml", wb)
        z.writestr("xl/_rels/workbook.xml.rels", wr)
        z.writestr("xl/worksheets/sheet1.xml", sh)
    return yol


def test_merge_bomba_tez_rad_etiladi(tmp_path):
    """1.4 KB fayl xotirani tugatmasin — sekundlar ichida rad etilsin."""
    bomba = _merge_bomba(str(tmp_path / "bomba.xlsx"))
    assert os.path.getsize(bomba) < 10_000, "hujum fayli kichkina bo'lishi kerak"
    t0 = time.time()
    with pytest.raises(ExcelTooLargeError):
        read_file(bomba)
    assert time.time() - t0 < 5, "rad etish sekin — himoya kech ishlayapti"


def test_oddiy_merge_hali_ishlaydi(tmp_path):
    """Haqiqiy hujjatlardagi kichik merge diapazonlari ochilishda davom etsin."""
    oddiy = _merge_bomba(str(tmp_path / "oddiy.xlsx"), ref="A1:D3")
    varaqlar = read_file(oddiy)
    assert len(varaqlar) == 1


@pytest.mark.parametrize("nom", [
    "etalon/1/1649137032-1. Жамланма жадвал.xlsx",
    "etalon/1/1649137090-Физобъемлар.xls",
    "etalon/2/1648641446-КАП_РЕМОНТ_№7_ШКОЛЫ_СЕРГЕЛИЙ_28,3,2022.xlsx",
])
def test_haqiqiy_etalonlar_oqiladi(nom):
    """Himoya haqiqiy fayllarni bloklamasligi kerak."""
    yol = os.path.join(ILDIZ, nom)
    if not os.path.exists(yol):
        pytest.skip(f"fayl yo'q: {nom}")
    varaqlar = read_file(yol)
    assert varaqlar, "varaqlar bo'sh"


# ═══════════════════════════════════════════════════════════════════════════
# Yuklab olish siyosati (blocker 3)
# ═══════════════════════════════════════════════════════════════════════════

class _Ishlovchi(BaseHTTPRequestHandler):
    kodlar = {}
    sarlavhalar = {}

    def log_message(self, *a):
        pass

    def do_GET(self):
        p = self.path.split("?")[0]
        _Ishlovchi.sarlavhalar = dict(self.headers)
        kod = _Ishlovchi.kodlar.get(p, 404)
        self.send_response(kod)
        self.send_header("Content-Length", "0")
        self.end_headers()


@pytest.fixture(scope="module")
def server():
    _Ishlovchi.kodlar = {"/401.xlsx": 401, "/403.xlsx": 403,
                         "/404.xlsx": 404, "/410.xlsx": 410, "/500.xlsx": 500}
    srv = HTTPServer(("127.0.0.1", 0), _Ishlovchi)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    eski = common.DOWNLOAD_RETRIES
    common.DOWNLOAD_RETRIES = 2
    # Sinov serveri 127.0.0.1 da — host oq ro'yxati (2026-09-23 da qo'shilgan)
    # uni rad etmasin. Bu testlar QAYTA URINISH siyosatini tekshiradi.
    eski_hostlar = common.RUXSAT_HOSTLAR
    common.RUXSAT_HOSTLAR = eski_hostlar + ("127.0.0.1",)
    yield f"http://127.0.0.1:{srv.server_address[1]}"
    common.DOWNLOAD_RETRIES = eski
    common.RUXSAT_HOSTLAR = eski_hostlar
    srv.shutdown()


@pytest.mark.parametrize("yol", ["/401.xlsx", "/403.xlsx", "/500.xlsx"])
def test_vaqtinchalik_xatolar_qayta_uriniladi(server, yol, tmp_path):
    """401/403/5xx — WAF yoki muddati o'tgan imzo bo'lishi mumkin.

    Bularni darhol "o'lik" qilish go-live'da butun navbatni yo'q qilardi.
    """
    with pytest.raises(Exception) as xato:
        common.download(server + yol, str(tmp_path))
    assert not isinstance(xato.value, common.PermanentDownloadError), \
        f"{yol} doimiy o'lim deb belgilandi — qayta urinilishi kerak edi"


@pytest.mark.parametrize("yol", ["/404.xlsx", "/410.xlsx"])
def test_umidsiz_xatolar_qayta_urinilmaydi(server, yol, tmp_path):
    """404/410 — havola haqiqatan yo'q, qayta urinish behuda."""
    with pytest.raises(common.PermanentDownloadError):
        common.download(server + yol, str(tmp_path))


def test_brauzerga_oxshash_user_agent(server, tmp_path):
    """`python-httpx` User-Agent ni ko'p WAF bloklaydi."""
    try:
        common.download(server + "/404.xlsx", str(tmp_path))
    except Exception:
        pass
    ua = _Ishlovchi.sarlavhalar.get("User-Agent", "")
    assert ua and "python-httpx" not in ua, f"shubhali User-Agent: {ua!r}"


def test_qoshimcha_sarlavhalar(server, tmp_path, monkeypatch):
    """DOWNLOAD_HEADERS orqali Authorization yuborish mumkin bo'lsin."""
    monkeypatch.setenv("DOWNLOAD_HEADERS", '{"Authorization": "Bearer SINOV"}')
    try:
        common.download(server + "/404.xlsx", str(tmp_path))
    except Exception:
        pass
    assert _Ishlovchi.sarlavhalar.get("Authorization") == "Bearer SINOV"


# ═══════════════════════════════════════════════════════════════════════════
# Deploy tayyorligi (blocker 1, 4, 6)
# ═══════════════════════════════════════════════════════════════════════════

@pytest.fixture(scope="module")
def toza_deploy(tmp_path_factory):
    """Faqat ishlab chiqarish fayllari solingan toza papka."""
    d = tmp_path_factory.mktemp("deploy")
    for f in PROD_FAYLLAR:
        manba = os.path.join(ILDIZ, f)
        if os.path.exists(manba):
            shutil.copy(manba, str(d / f))
    for paket in PROD_PAKETLAR:
        manba = os.path.join(ILDIZ, paket)
        if os.path.isdir(manba):
            shutil.copytree(manba, str(d / paket),
                            ignore=shutil.ignore_patterns("__pycache__"))
    return str(d)


def _ishlat(papka, argv, env_qosh=None, timeout=300):
    env = dict(os.environ)
    env.pop("DATABASE_URL", None)
    env.pop("JOBS_TABLE", None)
    env["PYTHONIOENCODING"] = "utf-8"
    if env_qosh:
        env.update(env_qosh)
    return subprocess.run([sys.executable] + argv, cwd=papka, env=env,
                          capture_output=True, text=True,
                          encoding="utf-8", errors="replace", timeout=timeout)


def test_prod_paketi_import_boladi(toza_deploy):
    """Deploy manifestidagi fayllar O'ZI yetarli bo'lsin — har kirish nuqtasi."""
    p = _ishlat(toza_deploy, ["-c", "import main, jobs_worker, yuboruvchi, api_server, "
                                    "korish, kuzatuv, ishga_tushir, app.db.migrate"])
    assert p.returncode == 0, f"import xatosi:\n{p.stderr[-600:]}"


def test_jadval_nomi_inyeksiyasi_rad_etiladi(toza_deploy):
    """Jadval nomi endi SOZLANMAYDI (sxema qat'iy) — boshqa qiymat, ayniqsa
    SQL bo'lagi, jim e'tiborsiz qolmasin: xizmat aniq xabar bilan to'xtaydi."""
    p = _ishlat(toza_deploy, ["-c", "import jobs_worker"],
                {"JOBS_TABLE": "jobs; DROP TABLE x--"})
    assert p.returncode != 0, "zararli jadval nomi qabul qilindi!"
    assert "JOBS_TABLE" in (p.stdout + p.stderr)


class _SoxtaKursor:
    """`SELECT count(*)` ga belgilangan sonni qaytaradigan soxta kursor."""

    def __init__(self, son):
        self.son = son

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def execute(self, *a, **k):
        pass

    def fetchone(self):
        return (self.son,)


class _SoxtaUlanish:
    def __init__(self, son):
        self.son = son

    def cursor(self):
        return _SoxtaKursor(self.son)

    def rollback(self):
        pass


def test_etalon_manbai_yoq_bolsa_toxtatadi(monkeypatch):
    """Sozlama XATOSI — o'lim; hali BO'SH jadval — kutish (2026-09-22).

    Etalon FAQAT `templates` jadvalidan olinadi (diskdagi zaxira 2026-09-29 da
    olib tashlandi).
      • jadval O'QILMASA (nomi noto'g'ri, huquq yo'q) → FATAL, systemd ko'rsin;
      • jadval BOR, lekin hali BO'SH → to'xtamaydi: yangi o'rnatishda sherik
        shablonlarni keyinroq yozadi, fayllar esa F1 bilan «shablon kutilmoqda»
        holatida saqlanadi. Ilgari bu ham FATAL edi va toza serverda konteyner
        cheksiz qayta ishga tushardi (Docker sinovida topildi).
    """
    from app.worker import loop as jw

    # 1) templates bo'sh → TO'XTAMAYDI (kutish holati)
    jw._etalon_manbaini_tekshir(_SoxtaUlanish(0))

    # 1b) jadval o'qilmasa (sozlama xatosi) → FATAL
    class _YiqiluvchiUlanish(_SoxtaUlanish):
        def cursor(self):
            raise RuntimeError("relation \"templates\" does not exist")

    with pytest.raises(SystemExit) as exc:
        jw._etalon_manbaini_tekshir(_YiqiluvchiUlanish(0))
    assert exc.value.code == 1

    # 2) templates to'la → ishlashda davom etadi.
    #    Ilgari shu holatda ham o'lardi — toza serverda worker ishga
    #    tushmasligining sababi aynan shu edi.
    jw._etalon_manbaini_tekshir(_SoxtaUlanish(5))


def test_nisbiy_yol_boshqa_cwd_dan_ishlaydi(toza_deploy):
    """CWD='/' (systemd, cron) bo'lsa ham nisbiy yo'llar LOYIHA ildiziga
    nisbatan hal qilinsin — ETALON_CACHE_DIR/DOWNLOAD_DIR boshqa joyga ketmasin."""
    ildiz = "C:\\" if os.name == "nt" else "/"
    env = dict(os.environ)
    env.pop("DATABASE_URL", None)
    env["PYTHONIOENCODING"] = "utf-8"
    env["ETALON_CACHE_DIR"] = "./_kesh_sinov"
    env["DOWNLOAD_DIR"] = "./_yuklash_sinov"
    kod = ("import sys; sys.path.insert(0, %r); from app import config; "
           "print(config.shablon().etalon_cache_dir); print(config.yuklash().download_dir)"
           % toza_deploy)
    p = subprocess.run([sys.executable, "-c", kod], cwd=ildiz, env=env,
                       capture_output=True, text=True, encoding="utf-8",
                       errors="replace", timeout=300)
    assert p.returncode == 0, p.stderr[-400:]
    kesh, yuklash = p.stdout.split()
    assert kesh == os.path.join(toza_deploy, "_kesh_sinov")
    assert yuklash == os.path.join(toza_deploy, "_yuklash_sinov")


def test_deploy_manifesti_toliq():
    """PROD_FAYLLAR ro'yxatidagi hamma fayl mavjudmi."""
    yetishmaydi = [f for f in PROD_FAYLLAR
                   if not os.path.exists(os.path.join(ILDIZ, f))]
    assert not yetishmaydi, f"deploy fayllari yo'q: {yetishmaydi}"


# ═══════════════════════════════════════════════════════════════════════════
# LOYIHA («5-ILOVA» / SMETA HISOBI) — 2026-07-31 da qo'shilgan `loyiha_excel`
# ═══════════════════════════════════════════════════════════════════════════

def _loyiha_fayli(yol, qiymatlar=None, futer_narx=1236000):
    """«SMETA HISOBI» shaklini yasaydi.

    `qiymatlar=None` → etalon (Qiymat ustuni bo'sh).
    Futer qatori («Jami: QQS bilan ...») HAR DOIM yoziladi — u merge qilingan
    matn bo'lib, o'ngida lotning boshlang'ich narxi turadi. Bu son
    ishtirokchining taklifi EMAS, shuning uchun «to'ldirilgan» deb
    sanalmasligi kerak.
    """
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "5-ILOVA"
    ws["A4"] = "SMETA HISOBI"
    for ustun, nom in zip("ABCD", ["T/r", "Ish turlari",
                                   "Narxning kelib chiqishi", "Qiymat (ming so‘m)"]):
        ws[f"{ustun}6"] = nom
        ws[f"{ustun}7"] = nom
    bandlar = ["1.", "2."]
    for i, band in enumerate(bandlar):
        r = 8 + i
        ws[f"A{r}"] = band
        ws[f"B{r}"] = f"Loyiha ishlari {i + 1}"
        ws[f"C{r}"] = "lot № 254110122"
        if qiymatlar:
            ws[f"D{r}"] = qiymatlar[i]
    ws["B11"] = "Jami QQS hisobga olmaganda:"
    ws["B12"] = "Jami QQS bilan:"
    if qiymatlar:
        ws["D11"] = sum(qiymatlar)
        ws["D12"] = sum(qiymatlar) * 1.12
    # futer: chapdagi matn merge qilingan, o'ngda lot narxi
    for ustun in "ABC":
        ws[f"{ustun}14"] = "Jami: QQS bilan ______ so‘m."
    ws["D14"] = futer_narx
    wb.save(yol)
    return yol


def test_loyiha_roli_aniqlanadi(tmp_path):
    """`loyiha_excel` shakli avtomatik tanilishi kerak."""
    from tender_engine.roles import detect_role
    from app import templates as templates_db
    yol = _loyiha_fayli(str(tmp_path / "etalon.xlsx"))
    assert detect_role(read_file(yol)) == "loyiha"
    # `type` bo'yicha zaxira ham bo'lsin — tuzilma o'zgarib ketsa ishlaydi
    assert templates_db.TYPE_ROL.get("loyiha_excel") == "loyiha"
    assert "loyiha" in templates_db.ROLLAR


def test_loyiha_toldirilgan_qabul_qilinadi(tmp_path):
    from tender_engine.validate import validate_one
    et = read_file(_loyiha_fayli(str(tmp_path / "e.xlsx")))
    pt = read_file(_loyiha_fayli(str(tmp_path / "p.xlsx"), [1048392, 500000]))
    n = validate_one("loyiha", et, pt, "p.xlsx")
    assert n["status"] == 1, n["comment_uz"]


def test_loyiha_bosh_rad_etiladi(tmp_path):
    """Futerdagi lot narxi hujjatni «to'ldirilgan» qilib ko'rsatmasin.

    Aynan shu tuzoq haqiqiy faylda uchradi: ishtirokchi bo'sh shaklni
    qaytargan, lekin futerda 1 236 000 turgani uchun qabul qilinib ketardi.
    """
    from tender_engine.validate import validate_one
    et = read_file(_loyiha_fayli(str(tmp_path / "e.xlsx")))
    pt = read_file(_loyiha_fayli(str(tmp_path / "p.xlsx")))   # Qiymat bo'sh
    n = validate_one("loyiha", et, pt, "p.xlsx")
    assert n["status"] == 2, "bo'sh hujjat qabul qilindi"
    assert "to'ldirilmagan" in n["comment_uz"], n["comment_uz"]


def test_loyiha_ustun_ochirilsa_ham_qabul(tmp_path):
    """Kelishilgan siyosat: FAQAT «Qiymat» muhim — ustun tuzilmasi emas."""
    from openpyxl import load_workbook
    from tender_engine.validate import validate_one
    et = read_file(_loyiha_fayli(str(tmp_path / "e.xlsx")))
    p = _loyiha_fayli(str(tmp_path / "p.xlsx"), [1048392, 500000])
    wb = load_workbook(p)
    wb.active.delete_cols(3)          # «Narxning kelib chiqishi» ni o'chiramiz
    wb.save(p)
    n = validate_one("loyiha", et, read_file(p), "p.xlsx")
    assert n["status"] == 1, f"ustun o'chirilgani uchun rad etildi: {n['comment_uz']}"


# ═══════════════════════════════════════════════════════════════════════════
# Q3 — «jiddiy to'ldirish» (2026-08-04): bir-ikkita qiymat yoki faqat JAMI
# yozib qo'yilgan hujjat RAD etilishi shart
# ═══════════════════════════════════════════════════════════════════════════

def _katta_loyiha_fayli(yol, qiymatlar=None, jami=None):
    """6 bandli «SMETA HISOBI» — Q3 sinovlari uchun.

    `qiymatlar` — har band uchun narx (None = bo'sh qoldiriladi).
    `jami` — «Jami QQS bilan» qatoriga yoziladigan son (faqat-jami holati).
    """
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "5-ILOVA"
    ws["A4"] = "SMETA HISOBI"
    for ustun, nom in zip("ABCD", ["T/r", "Ish turlari",
                                   "Narxning kelib chiqishi", "Qiymat (ming so‘m)"]):
        ws[f"{ustun}6"] = nom
    for i in range(6):
        r = 7 + i
        ws[f"A{r}"] = f"{i + 1}."
        ws[f"B{r}"] = f"Qurilish-montaj ishlari {i + 1}"
        if qiymatlar and qiymatlar[i] is not None:
            ws[f"D{r}"] = qiymatlar[i]
    ws["B14"] = "Jami QQS bilan:"
    if jami is not None:
        ws["D14"] = jami
    wb.save(yol)
    return read_file(yol)


def test_q3_bir_ikkita_qiymat_rad_etiladi(tmp_path):
    """6 banddan atigi 2 tasi to'ldirilgan — jiddiy emas, RAD."""
    from tender_engine.validate import validate_one
    et = _katta_loyiha_fayli(str(tmp_path / "e.xlsx"))
    pt = _katta_loyiha_fayli(str(tmp_path / "p.xlsx"),
                             [100, 200, None, None, None, None])
    n = validate_one("loyiha", et, pt, "p.xlsx")
    assert n["status"] == 2, f"2/6 qiymat bilan qabul qilindi: {n['comment_uz']}"
    assert any(f["code"] == "PRICE_SPARSE" for f in n["findings"]), n["findings"]


def test_q3_faqat_jami_rad_etiladi(tmp_path):
    """Bandlar bo'sh, faqat «Jami» qatoriga son yozilgan — RAD."""
    from tender_engine.validate import validate_one
    et = _katta_loyiha_fayli(str(tmp_path / "e.xlsx"))
    pt = _katta_loyiha_fayli(str(tmp_path / "p.xlsx"), jami=99_000_000)
    n = validate_one("loyiha", et, pt, "p.xlsx")
    assert n["status"] == 2, f"faqat-jami hujjat qabul qilindi: {n['comment_uz']}"
    assert any(f["code"] in ("PRICE_ONLY_TOTAL", "PRICE_EMPTY")
               for f in n["findings"]), n["findings"]


def test_q3_toliq_toldirilgan_qabul(tmp_path):
    """Hamma band to'ldirilgan — Q3 xalaqit bermasin."""
    from tender_engine.validate import validate_one
    et = _katta_loyiha_fayli(str(tmp_path / "e.xlsx"))
    pt = _katta_loyiha_fayli(str(tmp_path / "p.xlsx"),
                             [100, 200, 300, 400, 500, 600], jami=2100)
    n = validate_one("loyiha", et, pt, "p.xlsx")
    assert n["status"] == 1, n["comment_uz"]


def test_dublikat_sarlavhada_eng_boy_ustun_tanlanadi(tmp_path):
    """Ikkita «Qiymat» ustuni: birinchisi matnli izoh, haqiqiysi keyingisi.

    Haqiqiy bazada shu holat 10 ta to'ldirilgan faylni noo'rin rad ettirgan.
    """
    from openpyxl import Workbook
    from tender_engine.validate import validate_one

    def fayl(yol, toldirilganmi):
        wb = Workbook()
        ws = wb.active
        ws.title = "5-ILOVA"
        ws["A4"] = "SMETA HISOBI"
        nomlar = ["T/r", "Ish turlari", "Asos", "Qiymat (ming so‘m)",
                  "Narxning kelib chiqishi", "Qiymat (ming so‘m)"]
        for c, nom in enumerate(nomlar, start=1):
            ws.cell(row=6, column=c, value=nom)
        for i in range(4):
            r = 7 + i
            ws.cell(row=r, column=1, value=f"{i + 1}.")
            ws.cell(row=r, column=2, value=f"Ta'mirlash ishlari {i + 1}")
            # [3]-ustun («Qiymat») — MATNLI izoh, hamma faylda bor
            ws.cell(row=r, column=4, value="Obyektning taxminiy QMI qiymati")
            if toldirilganmi:
                ws.cell(row=r, column=6, value=1000000 + i)
        wb.save(yol)
        return read_file(yol)

    et = fayl(str(tmp_path / "e.xlsx"), False)
    pt = fayl(str(tmp_path / "p.xlsx"), True)
    n = validate_one("loyiha", et, pt, "p.xlsx")
    assert n["status"] == 1, f"to'ldirilgan fayl rad etildi: {n['comment_uz']}"
    bo_sh = validate_one("loyiha", et, fayl(str(tmp_path / "b.xlsx"), False), "b.xlsx")
    assert bo_sh["status"] == 2, "bo'sh fayl qabul qilindi"


def test_jamlanma_narx_langari_bilan_oqiladi(tmp_path):
    """D+ (7-taklif): tavsif ustuni nomsiz Форма-4 — narx nomi langar bo'ladi.

    MUHIM: bo'sh hujjat BARIBIR rad etilishi shart (name_idx == price_idx
    tuzog'i yopilganini shu tasdiqlaydi).
    """
    from openpyxl import Workbook
    from tender_engine.validate import validate_one

    def fayl(yol, qiymatlar=None):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "ФОРМА 4"
        ws["A3"] = "СВОДНАЯ ТАБЛИЦА"
        ws["B5"] = "наименование"          # tavsif ustunining «yarim» nomi
        ws["C5"] = "Стоимость в текущих ценах"
        ws["A6"], ws["B6"], ws["C6"] = 1, 2, 3        # raqamlash qatori
        bandlar = ["Затраты на оборудование", "Строительные работы",
                   "Монтажные работы", "Прочие затраты"]
        for i, b in enumerate(bandlar):
            r = 7 + i
            ws[f"A{r}"] = i + 1
            ws[f"B{r}"] = b
            if qiymatlar:
                ws[f"C{r}"] = qiymatlar[i]
        ws["B12"] = "ИТОГО:"
        if qiymatlar:
            ws["D12"] = sum(qiymatlar)
        wb.save(yol)
        return read_file(yol)

    from tender_engine.validate import _etalon_oqiladimi
    et = fayl(str(tmp_path / "e.xlsx"))
    assert _etalon_oqiladimi("jamlanma", et), "D+ langari etalonni ochmadi"

    pt = fayl(str(tmp_path / "p.xlsx"), [1000, 2000, 3000, 4000])
    assert validate_one("jamlanma", et, pt, "p.xlsx")["status"] == 1

    bosh = fayl(str(tmp_path / "b.xlsx"))
    n = validate_one("jamlanma", et, bosh, "b.xlsx")
    assert n["status"] == 2, "BO'SH hujjat narx-langar orqali qabul qilindi!"


# ═══════════════════════════════════════════════════════════════════════════
# «Tushunmadim» ≠ «rad etaman» (OLTIN QOIDA)
# ═══════════════════════════════════════════════════════════════════════════

def test_notanish_shabl_texnik_holat(tmp_path):
    """Buyurtmachi shabloni bizga notanish ko'rinishda bo'lsa — status 0.

    Haqiqiy bazada shunday etalonlar bor (o'zbekcha-lotin «UMUMIY JADVAL»).
    Ilgari `HEADER_NOT_FOUND` bilan status=2 chiqardi — ya'ni ishtirokchi
    BIZNING cheklovimiz uchun rad etilardi.
    """
    from openpyxl import Workbook
    from tender_engine.validate import validate_one, STATUS_TECHNICAL

    # 2026-08-04: fikstura nomlari yangilandi — «Ishlarning nomi» va
    # «To'g'ridan-to'g...» endi TANISH kalitlar (87 notanish shablon
    # tahlilidan qo'shildi), shuning uchun chindan notanish nomlar olindi.
    def notanish(yol, qiymat=None):
        wb = Workbook(); ws = wb.active
        ws["A6"] = "UMUMIY JADVAL"
        for u, n in zip("ABC", ["№ T.R", "Obyekt bo'limlari", "Mablag' miqdori"]):
            ws[f"{u}11"] = n
        ws["A12"] = 1
        ws["B12"] = "Qurilish ishlari"
        if qiymat:
            ws["C12"] = qiymat
        wb.save(yol)
        return read_file(yol)

    et = notanish(str(tmp_path / "e.xlsx"))
    pt = notanish(str(tmp_path / "p.xlsx"), 5000)
    n = validate_one("jamlanma", et, pt, "p.xlsx")
    assert n["status"] == STATUS_TECHNICAL, \
        f"notanish shablon uchun status={n['status']} berildi (0 kutilgan)"
    assert n["comment_uz"] == "", "texnik holatda izoh yozilmasin"
    assert any(f.get("severity") == "technical" for f in n["findings"])


def test_tanish_shabl_hamon_tekshiriladi(tmp_path):
    """Yangi himoya tanish shakllarni to'sib qo'ymasin."""
    from tender_engine.validate import validate_one
    et = read_file(_loyiha_fayli(str(tmp_path / "e.xlsx")))
    pt = read_file(_loyiha_fayli(str(tmp_path / "p.xlsx")))     # bo'sh
    assert validate_one("loyiha", et, pt, "p.xlsx")["status"] == 2
    pt2 = read_file(_loyiha_fayli(str(tmp_path / "p2.xlsx"), [100, 200]))
    assert validate_one("loyiha", et, pt2, "p2.xlsx")["status"] == 1


# ═══════════════════════════════════════════════════════════════════════════
# Deploy manifesti — Dockerfile shu ro'yxatni ko'chirsin
# ═══════════════════════════════════════════════════════════════════════════

def test_dockerfile_prod_fayllarni_oz_ichiga_oladi():
    """Dockerfile COPY ro'yxati PROD_FAYLLAR/PROD_PAKETLAR bilan mos bo'lsin.

    Aynan shunday desinxronlik tufayli `templates_db.py` bir paytlar
    Dockerfile ga tushmay qolgan va konteyner `ModuleNotFoundError` bilan
    yiqilardi. Endi paketlar butunicha ko'chiriladi — ro'yxat qisqa.
    """
    matn = open(os.path.join(ILDIZ, "Dockerfile"), encoding="utf-8").read()
    kerak = [f for f in PROD_FAYLLAR] + [f"{p}/" for p in PROD_PAKETLAR]
    yoq = [f for f in kerak if f not in matn]
    assert not yoq, f"Dockerfile COPY da yo'q: {yoq}"


def test_kirish_skriptlari_ingichka():
    """Ildizdagi kirish skriptlari FAQAT paketni chaqiradi — kod `app/` da."""
    for f in PROD_FAYLLAR:
        if not f.endswith(".py"):
            continue
        matn = open(os.path.join(ILDIZ, f), encoding="utf-8").read()
        assert "from app." in matn and len(matn.splitlines()) <= 15, f


#: `reader.read_file` ning zaxira zanjiri shu modullarga tayanadi. Ular
#: `tender_engine/` ichida, ya'ni PROD_PAKETLAR bilan ko'chiriladi — LEKIN
#: zanjir ularni KECH (xato yo'lida, `except` ichida) import qiladi, shuning
#: uchun yetishmagani konteyner qurilishida SEZILMAY qolardi va faqat
#: ishlab chiqarishda, buzuq fayl kelganda bilinardi.
ZAXIRA_MODULLARI = ["tiklash_rc4", "tiklash_xom_xml",
                    "tiklash_nol_bayt", "tiklash_xlsb"]


def test_zaxira_modullari_deploy_paketida_bor(toza_deploy):
    """Tiklash modullari toza deploy papkasida IMPORT bo'lsin."""
    kod = "import " + ", ".join("tender_engine.%s" % m for m in ZAXIRA_MODULLARI)
    p = _ishlat(toza_deploy, ["-c", kod])
    assert p.returncode == 0, f"zaxira moduli deploy paketida yo'q:\n{p.stderr[-600:]}"


def test_zaxira_modullari_manba_papkada_bor():
    """Fayllarning o'zi `tender_engine/` da turibdimi (arzon, tez tekshiruv)."""
    yoq = [m for m in ZAXIRA_MODULLARI
           if not os.path.isfile(os.path.join(ILDIZ, "tender_engine", m + ".py"))]
    assert not yoq, f"tender_engine/ da yo'q: {yoq}"


def test_compose_postgres_sozlamasi():
    """Compose dagi o'z PostgreSQL bazamiz XAVFSIZ sozlangan bo'lsin.

    Xavf o'zgarmagan: konteyner tasodifan BO'SH bazaga ulanib qolmasin,
    ma'lumot restartda yo'qolmasin, baza tashqariga ochilmasin. SHART:
      • ma'lumot nomli volume'da (`pgdata`) — konteyner o'chsa ham qoladi;
      • sxema migratsiya bilan yaratiladi (`migrate` xizmati), ilova esa u
        muvaffaqiyatli tugagandan KEYIN ishga tushadi;
      • migratsiya baza SOG'LOM bo'lgandan keyin yuradi;
      • bazaning porti tashqariga OCHILMAYDI (faqat ichki tarmoq).
    """
    matn = open(os.path.join(ILDIZ, "docker-compose.yml"), encoding="utf-8").read()
    assert "image: postgres" in matn, "compose da o'z bazamiz yo'q"
    assert "pgdata:/var/lib/postgresql/data" in matn, \
        "postgres ma'lumoti nomli volume'ga ulanmagan — konteyner qayta yaratilsa YO'QOLADI"
    assert '"-m", "app.db.migrate"' in matn, "migratsiya xizmati yo'q (jadvallar yaratilmaydi)"
    assert "condition: service_completed_successfully" in matn, \
        "ilova migratsiya tugashini kutmaydi"
    assert "condition: service_healthy" in matn, "migratsiya baza tayyor bo'lishini kutmaydi"
    blok = matn.split("\n  db:\n", 1)[1].split("\n  migrate:", 1)[0]
    assert "ports:" not in blok, "baza porti tashqariga ochilgan"
    assert "internal: true" in matn, "baza ichki tarmoqda emas"


def test_compose_da_qotirilgan_parol_yoq():
    """Compose faylida ochiq parol bo'lmasin — qiymatlar `${...}` orqali `.env` dan."""
    import re
    matn = open(os.path.join(ILDIZ, "docker-compose.yml"), encoding="utf-8").read()
    # 1) `POSTGRES_PASSWORD: <qiymat>` (YAML kaliti) faqat ${...} bo'lsin
    for m in re.finditer(r"^\s*POSTGRES_PASSWORD:\s*(\S.*)$", matn, re.M):
        qiymat = m.group(1).lstrip()
        # `""` — ilova konteynerlarida egasining parolini ATAYLAB bo'shatish
        assert qiymat.startswith("${") or qiymat.startswith('""'), \
            f"compose da ochiq parol: {m.group(0).strip()[:60]}"
    # 2) DSN da parol qotirilmagan bo'lsin (`postgresql://user:parol@...`)
    assert not re.search(r"postgresql://[^\s$'\"]+:(?!\$)[^\s@'\"]+@", matn), \
        "compose da parolli DSN qotirilgan"


def test_dockerignore_sirlarni_bloklaydi():
    matn = open(os.path.join(ILDIZ, ".dockerignore"), encoding="utf-8").read()
    for naqsh in (".env", ".venv/", "data_new/", "tests/"):
        assert naqsh in matn, f".dockerignore da yo'q: {naqsh}"
