# -*- coding: utf-8 -*-
"""KUZATUV — API integratsiyasi jonli analizi. FAQAT O'QIYDI (bazaga yozmaydi).

Savol: **qaysi fayl → qancha vaqtda → qanday verdikt → tenderga yetkazildimi.**

    python kuzatuv.py                 # oxirgi 24 soat: xulosa + vaqtlar + muammolar
    python kuzatuv.py --soat 2        # oxirgi 2 soat
    python kuzatuv.py --hammasi       # butun baza
    python kuzatuv.py --fayl 123      # BITTA faylning to'liq yo'li (timeline)
    python kuzatuv.py --file-id 4567  # tender tomondagi file_id bo'yicha
    python kuzatuv.py --muammo        # faqat muammolar (kechikkan, yuborilmagan, xato)
    python kuzatuv.py --kuzat         # jonli: har 15 s da yangilanadi (Ctrl+C)
    python kuzatuv.py --xlsx          # Excel hisobot (kuzatuv_<sana>.xlsx)
    python kuzatuv.py --chegara 50000 # bir so'rovda ko'proq qator (standart 5000)

Serverda: `docker compose exec worker python kuzatuv.py …`

Vaqt zanjiri (har fayl uchun):
    files.created_at ──► jobs_state.validated_at ──► yuborish_navbati.created_at ──► .yuborilgan_at
      API qabul qildi      tekshiruv tugadi           outbox'ga tushdi             tenderga yetdi
"""

import argparse
import collections
import os
import sys
import time

from app import config
from app.db import DATABASE_URL, schema
from app.db.schema import JOBS_TABLE, _yashamayotgan

#: Excel hisobot shu papkaga yoziladi (loyiha ildizi — avvalgidek).
BU = config.ROOT


VERDIKT = {0: "navbatda", 1: "QABUL", 2: "RAD", 3: "AYNAN NUSXA", 4: "TEXNIK(taslim)"}
HOLAT = {0: "navbatda", 1: "xato→qayta", 2: "YETKAZILDI", 3: "tashlandi"}

# Bitta so'rovda olinadigan eng ko'p qator. Sherik minglab fayl yuborganda
# chegara JIM qirqmasin — qirqilgani har chiqishda ochiq aytiladi va
# `--chegara` bilan ko'tariladi (`_ogohlantir`).
CHEGARA = 5000


def _ulan():
    import psycopg
    return psycopg.connect(DATABASE_URL, connect_timeout=60)


def _sek(a, b):
    return None if (a is None or b is None) else (b - a).total_seconds()


def _t(s):
    """Soniyani odam o'qiydigan ko'rinishga."""
    if s is None:
        return "—"
    if s < 60:
        return f"{s:.1f}s"
    if s < 3600:
        return f"{s / 60:.1f}daq"
    return f"{s / 3600:.1f}soat"


def _p(xs, p):
    xs = sorted(x for x in xs if x is not None)
    return xs[min(len(xs) - 1, int(p * len(xs)))] if xs else None


def _qatorlar(conn, soat=None, fayl=None, file_id=None):
    """Har fayl uchun butun zanjir bitta so'rovda (oxirgi outbox xabari bilan).

    Returns: (qatorlar, jami) — `jami` shartga mos QATORLARNING HAQIQIY soni.
    `len(qatorlar) < jami` bo'lsa chegara qirqqan (chaqiruvchi ogohlantiradi).

    Ustun nomlari ATAYLAB alias bilan (`ob_*`) — `files.created_at` va
    `yuborish_navbati.created_at` bir xil nom bo'lib, dict da bosib ketardi.
    """
    F = JOBS_TABLE
    sh, par = "TRUE", []
    if fayl:
        sh, par = "f.id = %s", [fayl]
    elif file_id:
        sh, par = "f.file_id = %s", [file_id]
    elif soat:
        sh, par = "f.created_at > now() - make_interval(hours => %s)", [soat]
    with conn.cursor() as cur:
        cur.execute(f"""
            -- `files.created_at` sherik sxemasida `timestamp(0)` (zonasiz), bizning
            -- jadvallar `timestamptz` — ayirishda to'qnashadi. `::timestamptz` uni
            -- server zonasi bo'yicha o'giradi (u `now()` bilan yozilgan).
            SELECT f.id, f.file_id, f.tender_id, f.type, f.link, f.status, f.comment,
                   f.created_at::timestamptz AS created_at,
                   f.updated_at::timestamptz AS updated_at,
                   s.validated_at, s.attempts, s.dead_at, s.last_error,
                   y.id            AS ob_id,
                   y.status        AS ob_status,
                   y.sabab         AS ob_sabab,
                   y.holat         AS ob_holat,
                   y.urinish       AS ob_urinish,
                   y.created_at    AS ob_yaratildi,
                   y.yuborilgan_at AS ob_yuborilgan,
                   y.xato          AS ob_xato
            FROM {F} f
            LEFT JOIN jobs_state s ON s.uuid = f.id::text
            LEFT JOIN LATERAL (
                SELECT * FROM yuborish_navbati q
                WHERE q.fayl_id = f.id ORDER BY q.id DESC LIMIT 1
            ) y ON TRUE
            WHERE {sh}{_yashamayotgan('f')}
            ORDER BY f.created_at DESC NULLS LAST, f.id DESC
            LIMIT {int(CHEGARA)}""", par)
        ust = [d[0] for d in cur.description]
        qatorlar = [dict(zip(ust, r)) for r in cur.fetchall()]
        if len(qatorlar) < CHEGARA:
            return qatorlar, len(qatorlar)
        # Chegaraga tegdik — haqiqiy sonni alohida so'raymiz (arzon, indeksli).
        cur.execute(f"SELECT count(*) FROM {F} f "
                    f"WHERE {sh}{_yashamayotgan('f')}", par)
        return qatorlar, cur.fetchone()[0]


def _ogohlantir(nechta, jami):
    """Chegara qirqqan bo'lsa OCHIQ aytamiz — jim qirqish chalg'itadi."""
    if jami > nechta:
        print(f"\n⚠️  DIQQAT: shartga {jami} ta fayl mos keladi, bu yerda faqat "
              f"ENG YANGI {nechta} tasi. To'liq ko'rish: --chegara {jami}")


def _muammo(x, hozir):
    """Fayl bo'yicha muammo bormi → sabab matni yoki None (yagona manba)."""
    if x["status"] == 0 and not x["validated_at"]:
        yosh = _sek(x["created_at"], hozir)
        if yosh and yosh > 900:                      # 15 daqiqa
            return f"tekshirilmagan {_t(yosh)} (urinish {x['attempts'] or 0})"
    elif x["status"] in (1, 2, 3, 4):
        if not x["ob_id"]:
            return "verdikt bor, outbox'da xabar YO'Q"
        if x["ob_holat"] == 3:
            return "tashlandi (yuborilmadi)"
        if x["ob_holat"] in (0, 1) and not x["ob_yuborilgan"]:
            yosh = _sek(x["ob_yaratildi"], hozir)
            if yosh and yosh > 600:                  # 10 daqiqa (cron 3 daq)
                return f"yuborilmagan {_t(yosh)} (urinish {x['ob_urinish']})"
    return None


def xulosa(conn, soat):
    import datetime
    q, jami_baza = _qatorlar(conn, soat=soat)
    if not q:
        print(f"Oxirgi {soat} soatda yangi fayl yo'q.")
        return 0
    hozir = datetime.datetime.now(datetime.timezone.utc)
    jami = len(q)
    verdikt = collections.Counter(x["status"] for x in q)
    yetk = collections.Counter()
    kutish, chiqish, toliq = [], [], []
    muammolar = []
    for x in q:
        yetk[x["ob_holat"] if x["ob_id"] else "outbox yo'q"] += 1
        kutish.append(_sek(x["created_at"], x["validated_at"]))
        chiqish.append(_sek(x["validated_at"], x["ob_yuborilgan"]))
        toliq.append(_sek(x["created_at"], x["ob_yuborilgan"]))
        s = _muammo(x, hozir)
        if s:
            muammolar.append((x, s, x["last_error"] or x["ob_xato"]))

    print("═" * 100)
    print(f"  KUZATUV — oxirgi {soat} soat · {jami} fayl · {time.strftime('%Y-%m-%d %H:%M')}")
    print("═" * 100)
    print("\n1) VERDIKT taqsimoti")
    for s, n in sorted(verdikt.items(), key=lambda t: (t[0] is None, t[0])):
        print(f"     {VERDIKT.get(s, s):16} {n:>6}   {100 * n / jami:5.1f}%")
    print("\n2) TENDERGA YETKAZISH")
    for h, n in sorted(yetk.items(), key=lambda t: str(t[0])):
        print(f"     {HOLAT.get(h, h):16} {n:>6}   {100 * n / jami:5.1f}%")
    print("\n3) VAQTLAR (p50 / p95 / eng uzun)")
    for nom, xs in (("qabul → tekshiruv", kutish), ("tekshiruv → yetkazish", chiqish),
                    ("JAMI (qabul → tender)", toliq)):
        n = sum(1 for x in xs if x is not None)
        print(f"     {nom:24} {_t(_p(xs, .5)):>8} / {_t(_p(xs, .95)):>8} / {_t(max([x for x in xs if x is not None], default=None) if n else None):>8}   ({n} ta)")
    if muammolar:
        print(f"\n4) MUAMMOLAR — {len(muammolar)} ta")
        for x, sabab, xato in muammolar[:15]:
            print(f"     id={x['id']:<8} file_id={x['file_id']:<10} tender={x['tender_id']:<10} "
                  f"{x['type']:<12} {sabab}")
            if xato:
                print(f"        └─ {str(xato)[:110]}")
        if len(muammolar) > 15:
            print(f"     … yana {len(muammolar) - 15} ta (--muammo bilan hammasi)")
    else:
        print("\n4) MUAMMOLAR — yo'q ✓")
    _ogohlantir(jami, jami_baza)
    return len(muammolar)


def bitta(conn, fayl=None, file_id=None):
    q, _ = _qatorlar(conn, fayl=fayl, file_id=file_id)
    if not q:
        print("Bunday fayl topilmadi.")
        return 1
    x = q[0]
    with conn.cursor() as cur:
        cur.execute("""SELECT id, status, sabab, holat, urinish, created_at, yuborilgan_at, xato, comment
                       FROM yuborish_navbati WHERE fayl_id = %s ORDER BY id""", (x["id"],))
        xabarlar = cur.fetchall()
        cur.execute("""SELECT ichki_status, tashqi_status, kodlar, comment_uz, yaratildi, validator_version
                       FROM validation_evidence WHERE fayl_id = %s ORDER BY id""", (str(x["id"]),))
        izlar = cur.fetchall()
    print("═" * 100)
    print(f"  FAYL  id={x['id']}  file_id={x['file_id']}  tender_id={x['tender_id']}  type={x['type']}")
    print("═" * 100)
    print(f"  havola   : {x['link'] or '—'}")
    print(f"  VERDIKT  : {VERDIKT.get(x['status'], x['status'])}")
    print(f"  izoh     : {x['comment'] or '—'}")
    print("\n  YO'L:")
    print(f"    {x['created_at']}  API qabul qildi")
    if x["validated_at"]:
        print(f"    {x['validated_at']}  tekshiruv tugadi      (+{_t(_sek(x['created_at'], x['validated_at']))})")
    else:
        print(f"    {'—':26}  HALI TEKSHIRILMAGAN   (urinish: {x['attempts'] or 0}"
              + (f", oxirgi xato: {str(x['last_error'])[:60]}" if x["last_error"] else "") + ")")
    for (yid, st, sabab, holat, urinish, yar, yub, xato, kom) in xabarlar:
        print(f"    {yar}  outbox #{yid} ({sabab}, status={st})")
        if yub:
            print(f"    {yub}  TENDERGA YETKAZILDI   (+{_t(_sek(yar, yub))}, urinish {urinish})")
        else:
            print(f"    {'—':26}  {HOLAT.get(holat, holat)} (urinish {urinish})"
                  + (f" — {str(xato)[:60]}" if xato else ""))
    if izlar:
        print("\n  AUDIT IZI (validation_evidence):")
        for ichki, tashqi, kodlar, kom, vaqt, ver in izlar:
            print(f"    {vaqt}  {ichki}/{tashqi}  {kodlar}  {ver}")
            if kom:
                print(f"        {str(kom)[:90]}")
    return 0


def muammolar_royxati(conn, soat):
    import datetime
    hozir = datetime.datetime.now(datetime.timezone.utc)
    chiqdi = 0
    q, jami_baza = _qatorlar(conn, soat=soat)
    for x in q:
        sabab = _muammo(x, hozir)
        if not sabab:
            continue
        chiqdi += 1
        print(f"id={x['id']:<8} file_id={x['file_id']:<10} tender={x['tender_id']:<10} "
              f"{x['type']:<12} status={VERDIKT.get(x['status'], x['status']):<14} {sabab}")
        xato = x["last_error"] or x["ob_xato"]
        if xato:
            print(f"    └─ {str(xato)[:120]}")
    print(f"\njami muammo: {chiqdi}")
    _ogohlantir(len(q), jami_baza)
    return chiqdi


def xlsx_hisobot(conn, soat):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    q, jami_baza = _qatorlar(conn, soat=soat)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kuzatuv"
    bosh = ["files.id", "file_id", "tender_id", "type", "verdikt", "izoh",
            "qabul (created_at)", "tekshiruv (validated_at)", "qabul→tekshiruv (s)",
            "outbox", "tenderga yetdi", "tekshiruv→yetkazish (s)", "JAMI (s)",
            "yetkazish holati", "urinish", "xato"]
    ws.append(bosh)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for x in q:
        ws.append([x["id"], x["file_id"], x["tender_id"], x["type"],
                   VERDIKT.get(x["status"], x["status"]), (x["comment"] or "")[:200],
                   _bz(x["created_at"]), _bz(x["validated_at"]),
                   _sek(x["created_at"], x["validated_at"]),
                   _bz(x["ob_yaratildi"]), _bz(x["ob_yuborilgan"]),
                   _sek(x["validated_at"], x["ob_yuborilgan"]),
                   _sek(x["created_at"], x["ob_yuborilgan"]),
                   HOLAT.get(x["ob_holat"], x["ob_holat"]), x["ob_urinish"],
                   (str(x["ob_xato"])[:200] if x["ob_xato"] else "")])
    ws.freeze_panes = "A2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(40, max(
            10, max(len(str(c.value or "")) for c in col[:200]) + 2))
    yol = os.path.join(BU, f"kuzatuv_{time.strftime('%Y-%m-%d_%H%M')}.xlsx")
    wb.save(yol)
    print(f"yozildi: {yol}  ({len(q)} qator)")
    _ogohlantir(len(q), jami_baza)
    return 0


def _bz(dt):
    """Excel uchun vaqt zonasiz (openpyxl tz-aware datetime ni qabul qilmaydi)."""
    return dt.replace(tzinfo=None) if dt is not None else None


def main():
    global CHEGARA                   # `--chegara` bilan ko'tariladi
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--soat", type=int, default=24)
    ap.add_argument("--hammasi", action="store_true")
    ap.add_argument("--fayl", type=int, help="files.id")
    ap.add_argument("--file-id", type=int, dest="file_id", help="tender tomondagi file_id")
    ap.add_argument("--muammo", action="store_true")
    ap.add_argument("--kuzat", action="store_true", help="jonli, har 15 s")
    ap.add_argument("--xlsx", action="store_true")
    ap.add_argument("--chegara", type=int, default=CHEGARA,
                    help=f"bir so'rovda eng ko'p qator (standart {CHEGARA})")
    a = ap.parse_args()
    CHEGARA = max(1, a.chegara)
    soat = 24 * 365 * 10 if a.hammasi else a.soat
    with _ulan() as conn:
        schema.tekshir(conn)
        if a.fayl or a.file_id:
            return bitta(conn, fayl=a.fayl, file_id=a.file_id)
        if a.muammo:
            return muammolar_royxati(conn, soat)
        if a.xlsx:
            return xlsx_hisobot(conn, soat)
        if a.kuzat:
            try:
                while True:
                    os.system("cls" if os.name == "nt" else "clear")
                    xulosa(conn, soat)
                    print("\n(Ctrl+C — to'xtatish; 15 s da yangilanadi)")
                    time.sleep(15)
            except KeyboardInterrupt:
                return 0
        return 0 if xulosa(conn, soat) == 0 else 0


if __name__ == "__main__":
    sys.exit(main())
